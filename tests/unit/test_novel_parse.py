"""Reading an uploaded novel: encodings, .docx, chapters and paragraphs (novel/parse.py)."""

from __future__ import annotations

import codecs
import io
from pathlib import Path

import pytest
from docx import Document

from drama_series_agent.core.paths import REPO_ROOT
from drama_series_agent.novel.parse import (
    MAX_UPLOAD_BYTES,
    PART_TARGET_CHARS,
    NovelParseError,
    ParsedNovel,
    clean_paragraphs,
    clean_title,
    compose,
    is_heading,
    parse_upload,
)

DEMO = REPO_ROOT / "tests" / "fixtures" / "demo" / "雨停之前.txt"
# The prototype's opening (web/src/mocks/fixtures/demo/novel.ts).
OPENING = [
    "那天雨下得很大。林晚加完班回到老小区，一层一层爬上六楼。",
    "家门口靠着一把红色的长柄伞，伞柄上贴着一张便利贴。她认得那个字迹。三年了，"
    "周屿的字还是歪歪扭扭的，像他这个人一样不靠谱。",
    "便利贴上只写了一句话：“楼下便利店，等你到十二点。”",
    "她把伞靠在墙边，进了屋，又出来，又进去。",
    "最后，她还是拿起了那把伞。",
]


def _paras(n: ParsedNovel, chapter: int) -> list[str]:
    return [n.text[s:e] for s, e in n.chapters[chapter].paras]


def _code(exc: pytest.ExceptionInfo[NovelParseError]) -> str:
    return exc.value.code


def test_demo_novel_has_three_chapters_and_the_prototype_opening() -> None:
    n = parse_upload(DEMO.name, DEMO.read_bytes())
    assert (n.encoding, n.split_by) == ("utf-8", "heading")
    assert [c.title for c in n.chapters] == ["第一章 红伞", "第二章 十二点", "第三章 旧发圈"]
    assert _paras(n, 0) == OPENING
    # Paragraphs are numbered across the book.
    assert [c.para_start for c in n.chapters] == [0, 5, 5 + len(n.chapters[1].paras)]
    assert n.chars == sum(len(p) for c in range(3) for p in _paras(n, c))
    assert 2000 < n.chars < 3500
    # A chapter runs from its heading to its last paragraph.
    first = n.chapters[0]
    assert n.text[first.start : first.end].startswith("第一章")
    assert n.text[first.start : first.end].endswith("最后，她还是拿起了那把伞。")


@pytest.mark.parametrize(
    ("data", "encoding"),
    [
        (codecs.BOM_UTF8 + "第一章　雨\n正文。".encode(), "utf-8"),
        ("第一章　雨\n正文。".encode(), "utf-8"),
        ("第一章　雨\n正文。".encode("gb18030"), "gb18030"),
        ("第一章　雨\n正文。".encode("utf-16"), "utf-16"),  # with a BOM
    ],
)
def test_encodings(data: bytes, encoding: str) -> None:
    n = parse_upload("a.txt", data)
    assert n.encoding == encoding
    assert n.text == "第一章　雨\n正文。"
    assert _paras(n, 0) == ["正文。"]


def test_unreadable_text() -> None:
    with pytest.raises(NovelParseError) as exc:
        parse_upload("a.txt", "Chapter 1\nRain.".encode("utf-16-le"))  # no BOM: NULs
    assert _code(exc) == "novel_unreadable"
    with pytest.raises(NovelParseError) as exc:
        parse_upload("a.txt", b"abc\x80def")  # neither UTF-8 nor GB18030
    assert _code(exc) == "novel_unreadable"


def test_line_ends_and_whitespace_are_normalized() -> None:
    n = parse_upload("a.TXT", "第1章 开始\r\n\r\n　　她来了。\r\n  他走了。 \r".encode())
    assert n.text == "第1章 开始\n\n　　她来了。\n  他走了。 \n"
    assert _paras(n, 0) == ["她来了。", "他走了。"]
    # Offsets point into the normalized text, without the indentation.
    s, _ = n.chapters[0].paras[0]
    assert n.text[s - 2 : s] == "　　"


@pytest.mark.parametrize(
    "line",
    ["第一章 红伞", "第一章红伞", "第 12 章", "第十二回　风雨", "第3节", "Chapter 3: Rain",
     "CHAPTER IV", "楔子", "番外一 初遇", "Prologue"],
)  # fmt: skip
def test_headings(line: str) -> None:
    assert is_heading(line)


@pytest.mark.parametrize(
    "line",
    ["第三节课下课后，他来了。", "第一次见他是在雨里。", "第二天，他就走了", "章节", "x" * 41,
     "第一章" + "很长" * 20],
)  # fmt: skip
def test_not_headings(line: str) -> None:
    assert not is_heading(line)


def test_text_before_the_first_heading_and_empty_headings() -> None:
    text = "雨停之前\n作者：某某\n\n第一卷\n第一章 红伞\n她来了。\n第二章 十二点\n他走了。\n"
    text += "第三章\n"
    n = parse_upload("a.txt", text.encode())
    # The front matter is a chapter without a title; 「第一卷」 is not a heading but a short
    # line, so it stays in the front matter; 「第三章」 has nothing under it and is dropped.
    assert [c.title for c in n.chapters] == [None, "第一章 红伞", "第二章 十二点"]
    assert _paras(n, 0) == ["雨停之前", "作者：某某", "第一卷"]
    assert [c.idx for c in n.chapters] == [0, 1, 2]


def test_heading_right_before_a_heading_is_dropped() -> None:
    n = parse_upload("a.txt", "第一回 序\n第一章 开始\n正文。\n".encode())
    assert [c.title for c in n.chapters] == ["第一章 开始"]


def test_no_headings_splits_by_length() -> None:
    para = "雨" * 1000
    n = parse_upload("a.txt", "\n".join([para] * 14).encode())
    assert n.split_by == "length"
    assert all(c.title is None for c in n.chapters)
    assert [len(c.paras) for c in n.chapters] == [6, 6, 2]
    assert all(c.chars <= PART_TARGET_CHARS for c in n.chapters)
    assert [c.para_start for c in n.chapters] == [0, 6, 12]


def test_docx_paragraphs_and_heading_styles() -> None:
    doc = Document()
    doc.add_heading("雨夜", level=1)
    doc.add_paragraph("那天雨下得很大。")
    doc.add_paragraph("第二章 十二点")
    doc.add_paragraph("她跑下楼。")
    buf = io.BytesIO()
    doc.save(buf)
    n = parse_upload("雨停之前.docx", buf.getvalue())
    assert n.encoding == "docx"
    # A Word heading style counts even without 「第 X 章」.
    assert [c.title for c in n.chapters] == ["雨夜", "第二章 十二点"]
    assert _paras(n, 0) == ["那天雨下得很大。"]


def test_broken_docx() -> None:
    with pytest.raises(NovelParseError) as exc:
        parse_upload("a.docx", b"not a zip")
    assert _code(exc) == "novel_unreadable"
    assert exc.value.params == {"reason": "docx"}


def test_refusals(tmp_path: Path) -> None:
    with pytest.raises(NovelParseError) as exc:
        parse_upload("a.pdf", b"%PDF")
    assert (_code(exc), exc.value.params) == ("unsupported_file", {"ext": ".pdf"})
    with pytest.raises(NovelParseError) as exc:
        parse_upload("a.txt", b" \n\n\t\n")
    assert _code(exc) == "novel_empty"
    with pytest.raises(NovelParseError) as exc:
        parse_upload("a.txt", b"x" * (MAX_UPLOAD_BYTES + 1))
    assert _code(exc) == "file_too_large"


def _layout(parsed: ParsedNovel) -> list[tuple[str | None, list[str]]]:
    return [(c.title, [parsed.text[a:b] for a, b in c.paras]) for c in parsed.chapters]


def test_compose_reimports_to_the_same_chapters() -> None:
    demo = parse_upload(DEMO.name, DEMO.read_bytes())
    edited = compose(_layout(demo), demo.split_by)
    assert edited.encoding == "edited" and edited.chars == demo.chars
    assert _layout(parse_upload("x.txt", edited.text.encode("utf-8"))) == _layout(demo)
    for c in edited.chapters:
        assert [edited.text[a:b] for a, b in c.paras] == _layout(demo)[c.idx][1]
        assert c.para_start == demo.chapters[c.idx].para_start


def test_compose_keeps_any_title_and_untitled_parts() -> None:
    edited = compose([(None, ["前言一段"]), ("红伞", ["一", "二"])], "heading")
    assert [(c.title, c.para_start, len(c.paras)) for c in edited.chapters] == [
        (None, 0, 1),
        ("红伞", 1, 2),
    ]
    assert edited.text == "前言一段\n\n红伞\n\n一\n二\n"


def test_cleaning_what_the_user_typed() -> None:
    assert clean_paragraphs(["  第一段\u3000", "第二段\r\n\n第三段", "", " \n "]) == [
        "第一段",
        "第二段",
        "第三段",
    ]
    assert clean_title("  第一章\u3000 红伞\n") == "第一章 红伞"
    assert clean_title("   ") is None and clean_title(None) is None
