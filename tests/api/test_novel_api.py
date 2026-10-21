"""``/series`` and ``/series/{id}/novel`` (SPEC 9.1): series list / create / rename, import."""

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import func, select

from drama_series_agent.api.app import create_app
from drama_series_agent.api.state import StudioState
from drama_series_agent.core.models_db import NovelSource
from drama_series_agent.core.paths import REPO_ROOT
from drama_series_agent.pipeline.drafts import DependencyDraft, VersionDraft
from drama_series_agent.pipeline.invalidation import slice_hash

LOCAL = "http://127.0.0.1"
V1 = "/api/v1"
DEMO = (REPO_ROOT / "tests" / "fixtures" / "demo" / "雨停之前.txt").read_bytes()
TITLES = ["第一章 红伞", "第二章 十二点", "第三章 旧发圈"]


@pytest.fixture
def client() -> Iterator[TestClient]:
    with TestClient(create_app(), base_url=LOCAL, headers={"X-Drama-Studio": "1"}) as c:
        yield c


def _state(client: TestClient) -> StudioState:
    state: StudioState = client.app.state.studio  # type: ignore[attr-defined]
    return state


def _code(resp: Any) -> str:
    return resp.json()["detail"]["code"]  # type: ignore[no-any-return]


def _upload(
    client: TestClient,
    series: str,
    data: bytes,
    name: str = "novel.txt",
    **form: str,
) -> Any:
    return client.post(
        f"{V1}/series/{series}/novel", files={"file": (name, data)}, data=form or None
    )


def _new_series(client: TestClient, title: str = "新剧") -> str:
    resp = client.post(f"{V1}/series", json={"title": title})
    assert resp.status_code == 201
    return str(resp.json()["id"])


def _novel_product(client: TestClient, series: str) -> dict[str, Any]:
    listed = client.get(f"{V1}/products", params={"owner": f"series:{series}"}).json()
    return next(p for p in listed if p["slot"] == "novel")


# ---------------------------------------------------------------------------- series


def test_list_create_rename(client: TestClient) -> None:
    assert [s["id"] for s in client.get(f"{V1}/series").json()] == ["demo"]
    sid = _new_series(client, "  Rain   Stops ")
    created = client.get(f"{V1}/series/{sid}").json()
    assert (created["title"], created["slug"], created["ui_lang"]) == (
        "Rain Stops",
        "rain-stops-2",  # the demo series has rain-stops
        "zh-CN",
    )
    # Newest first; the new series has the default output spec.
    assert [s["id"] for s in client.get(f"{V1}/series").json()] == [sid, "demo"]
    spec = client.get(f"{V1}/series/{sid}/spec").json()
    assert (spec["aspect"], spec["locked_at"]) == ("9:16", None)
    # Titles without ASCII words get a generic slug.
    other = client.get(f"{V1}/series/{_new_series(client, '雨停之前')}").json()
    assert other["slug"] == "series"

    renamed = client.patch(f"{V1}/series/{sid}", json={"title": "雨停之前 2", "ui_lang": "en"})
    assert (renamed.json()["title"], renamed.json()["ui_lang"]) == ("雨停之前 2", "en")
    assert renamed.json()["slug"] == "rain-stops-2"  # the slug does not follow the title


def test_series_errors(client: TestClient) -> None:
    assert client.post(f"{V1}/series", json={"title": "   "}).status_code == 422
    assert client.post(f"{V1}/series", json={"title": "x" * 81}).status_code == 422
    assert client.post(f"{V1}/series", json={"title": "a", "ui_lang": "fr"}).status_code == 422
    missing = client.patch(f"{V1}/series/nope", json={"title": "a"})
    assert (missing.status_code, _code(missing)) == (404, "series_not_found")


# ---------------------------------------------------------------------------- reading


def test_demo_novel_by_chapter(client: TestClient) -> None:
    novel = client.get(f"{V1}/series/demo/novel").json()
    assert novel["source"]["id"] == "demo-src"
    assert novel["source"]["filename"] == "雨停之前.txt"
    assert novel["source"]["rights_confirmed_at"] == "2026-10-14T01:22:00Z"
    assert [c["title"] for c in novel["chapters"]] == TITLES
    assert [c["id"] for c in novel["chapters"]] == ["demo-src-ch1", "demo-src-ch2", "demo-src-ch3"]
    assert novel["chapter"] == 0
    paras = novel["paragraphs"]
    assert [p["idx"] for p in paras] == [0, 1, 2, 3, 4]
    assert paras[0]["text"] == "那天雨下得很大。林晚加完班回到老小区，一层一层爬上六楼。"
    assert paras[-1]["text"] == "最后，她还是拿起了那把伞。"

    third = client.get(f"{V1}/series/demo/novel", params={"chapter": 2}).json()
    assert third["chapter"] == 2
    assert third["paragraphs"][0]["idx"] == novel["chapters"][2]["para_start"]
    assert third["paragraphs"][-1]["text"] == "红伞底下，雨还没有停。"
    # Offsets are the same text the paragraphs show.
    text = DEMO.decode()
    assert all(text[p["start"] : p["end"]] == p["text"] for p in third["paragraphs"])


def test_reading_errors_and_no_novel(client: TestClient) -> None:
    sid = _new_series(client)
    assert client.get(f"{V1}/series/{sid}/novel").json() == {
        "series_id": sid,
        "source": None,
        "chapters": [],
        "chapter": None,
        "paragraphs": [],
    }
    beyond = client.get(f"{V1}/series/demo/novel", params={"chapter": 3})
    assert (beyond.status_code, _code(beyond)) == (404, "chapter_not_found")
    assert client.get(f"{V1}/series/demo/novel", params={"chapter": -1}).status_code == 422
    missing = client.get(f"{V1}/series/nope/novel")
    assert (missing.status_code, _code(missing)) == (404, "series_not_found")


# ---------------------------------------------------------------------------- import


def test_dry_run_stores_nothing(client: TestClient) -> None:
    sid = _new_series(client)
    resp = _upload(client, sid, DEMO, "雨停之前.txt", dry_run="true")
    assert resp.status_code == 200
    found = resp.json()
    assert (found["dry_run"], found["source_id"], found["encoding"]) == (True, None, "utf-8")
    assert [c["title"] for c in found["chapters"]] == TITLES
    assert [c["id"] for c in found["chapters"]] == [None, None, None]
    assert client.get(f"{V1}/series/{sid}/novel").json()["source"] is None


def test_import_needs_the_rights_confirmation(client: TestClient) -> None:
    sid = _new_series(client)
    resp = _upload(client, sid, DEMO)
    assert (resp.status_code, _code(resp)) == (422, "rights_not_confirmed")
    assert client.get(f"{V1}/series/{sid}/novel").json()["source"] is None


def test_import_makes_an_adopted_import_version(client: TestClient) -> None:
    sid = _new_series(client)
    resp = _upload(client, sid, DEMO, "雨停之前.txt", rights_confirmed="true")
    assert resp.status_code == 200
    body = resp.json()
    assert body["dry_run"] is False and body["unchanged"] is False
    novel = client.get(f"{V1}/series/{sid}/novel").json()
    assert novel["source"]["id"] == body["source_id"]
    assert [c["id"] for c in novel["chapters"]] == [c["id"] for c in body["chapters"]]

    product = _novel_product(client, sid)
    assert (product["component_id"], product["status"]) == ("import", "adopted")
    (v1,) = product["versions"]
    assert (v1["origin"], v1["note"]) == ("import", "雨停之前.txt")
    assert product["adopted_version_id"] == v1["id"]
    assert v1["payload"]["source_id"] == body["source_id"]
    assert [c["idx"] for c in v1["payload"]["chapters"]] == [0, 1, 2]


def test_reimport_replaces_and_keeps_the_old_text(client: TestClient) -> None:
    changed = DEMO.replace("便利店里没有他。".encode(), "便利店里空无一人。".encode())
    first = _upload(client, "demo", changed, "v2.txt", rights_confirmed="true").json()
    assert first["unchanged"] is False
    novel = client.get(f"{V1}/series/demo/novel", params={"chapter": 1}).json()
    assert novel["source"]["id"] == first["source_id"] != "demo-src"
    assert novel["paragraphs"][-1]["text"] == "便利店里空无一人。"
    versions = _novel_product(client, "demo")["versions"]
    assert [(v["n"], v["origin"]) for v in versions] == [(1, "import"), (2, "import")]
    assert _novel_product(client, "demo")["adopted_version_id"] == versions[1]["id"]

    # The same text again adds nothing.
    again = _upload(client, "demo", changed, "same.txt", rights_confirmed="true").json()
    assert (again["unchanged"], again["source_id"]) == (True, first["source_id"])
    assert len(_novel_product(client, "demo")["versions"]) == 2

    async def sources() -> int:
        async with _state(client).db.session() as session:
            return (await session.exec(select(func.count()).select_from(NovelSource))).one()

    assert client.portal.call(sources) == 2  # type: ignore[union-attr]


def test_downstream_turns_stale_only_when_its_chapter_changes(client: TestClient) -> None:
    """A product built from chapter 1 (idx 0): stale after that chapter changes, not before."""
    state = _state(client)
    novel = _novel_product(client, "demo")
    chapter0 = novel["versions"][0]["payload"]["chapters"][0]

    async def downstream() -> str:
        p = await state.products.ensure(
            series_id="demo", owner_type="series", owner_id="demo", slot="test_downstream",
            component_id="test",
        )  # fmt: skip
        dep = DependencyDraft(novel["id"], slice_hash(chapter0), "chapter:0")
        draft = VersionDraft(payload={}, dependencies=[dep])
        await state.products.append(p.id, draft, origin="run")
        return p.id

    pid = client.portal.call(downstream)  # type: ignore[union-attr]

    def status() -> str:
        return str(client.get(f"{V1}/products/{pid}").json()["status"])

    assert status() == "adopted"
    later = DEMO.replace("便利店里没有他。".encode(), "便利店里空无一人。".encode())
    _upload(client, "demo", later, rights_confirmed="true")
    assert status() == "adopted"
    opening = later.replace("最后，她还是拿起了那把伞。".encode(), "最后，她拿起了伞。".encode())
    _upload(client, "demo", opening, rights_confirmed="true")
    assert status() == "stale"
    # Back to the first chapter it was built from: fresh again.
    _upload(client, "demo", DEMO, rights_confirmed="true")
    assert status() == "adopted"


def test_locked_novel_refuses_an_import(client: TestClient) -> None:
    pid = _novel_product(client, "demo")["id"]
    assert client.post(f"{V1}/products/{pid}/lock").status_code == 200
    changed = DEMO.replace("便利店里没有他。".encode(), "便利店里空无一人。".encode())
    resp = _upload(client, "demo", changed, rights_confirmed="true")
    assert (resp.status_code, _code(resp)) == (409, "product_locked")
    assert len(_novel_product(client, "demo")["versions"]) == 1


def test_docx_import(client: TestClient) -> None:
    doc = Document()
    doc.add_paragraph("第一章 红伞")
    doc.add_paragraph("那天雨下得很大。")
    buf = io.BytesIO()
    doc.save(buf)
    sid = _new_series(client)
    body = _upload(client, sid, buf.getvalue(), "雨停之前.docx", rights_confirmed="true").json()
    assert body["encoding"] == "docx"
    novel = client.get(f"{V1}/series/{sid}/novel").json()
    assert [p["text"] for p in novel["paragraphs"]] == ["那天雨下得很大。"]


@pytest.mark.parametrize(
    ("name", "data", "status", "code"),
    [
        ("a.pdf", b"%PDF", 415, "unsupported_file"),
        ("a.txt", b"  \n", 422, "novel_empty"),
        ("a.docx", b"not a zip", 422, "novel_unreadable"),
        ("a.txt", b"abc\x80def", 422, "novel_unreadable"),
    ],
)
def test_import_errors(client: TestClient, name: str, data: bytes, status: int, code: str) -> None:
    sid = _new_series(client)
    resp = _upload(client, sid, data, name, rights_confirmed="true")
    assert (resp.status_code, _code(resp)) == (status, code)
    missing = _upload(client, "nope", DEMO, rights_confirmed="true")
    assert (missing.status_code, _code(missing)) == (404, "series_not_found")


# ---------------------------------------------------------------------------- 逐段编辑


def _chapter(client: TestClient, idx: int, series: str = "demo") -> dict[str, Any]:
    data: dict[str, Any] = client.get(f"{V1}/series/{series}/novel", params={"chapter": idx}).json()
    return data


def _edit(client: TestClient, idx: int, title: str | None, paragraphs: list[str]) -> Any:
    body = {"title": title, "paragraphs": paragraphs}
    return client.put(f"{V1}/series/demo/novel/chapters/{idx}", json=body)


def _texts(chapter: dict[str, Any]) -> list[str]:
    return [p["text"] for p in chapter["paragraphs"]]


def test_editing_a_paragraph_makes_an_adopted_edit_version(client: TestClient) -> None:
    before = _chapter(client, 0)
    third = _texts(_chapter(client, 2))
    paras = _texts(before)
    paras[1] = paras[1].replace("歪歪扭扭", "七扭八歪")
    resp = _edit(client, 0, before["chapters"][0]["title"], paras)
    assert resp.status_code == 200, resp.text
    after = resp.json()
    assert after["source"]["id"] != before["source"]["id"]
    assert after["source"]["encoding"] == "edited"
    assert after["source"]["filename"] == before["source"]["filename"]
    assert _texts(after) == paras and [p["idx"] for p in after["paragraphs"]] == [0, 1, 2, 3, 4]
    # Offsets point into the new text.
    text = after["paragraphs"][1]
    assert text["end"] - text["start"] == len(text["text"])
    product = _novel_product(client, "demo")
    edit = product["versions"][-1]
    assert (edit["origin"], edit["parent_version_id"]) == ("edit", "demo-novel-v1")
    assert product["adopted_version_id"] == product["current_version_id"] == edit["id"]
    # The other chapters read as before; the same number of characters (歪歪扭扭 → 七扭八歪).
    assert _texts(_chapter(client, 2)) == third
    assert after["source"]["chars"] == before["source"]["chars"]
    # Same text again: nothing new.
    assert _edit(client, 0, before["chapters"][0]["title"], paras).status_code == 200
    assert len(_novel_product(client, "demo")["versions"]) == 2


def test_inserting_and_deleting_paragraphs_renumbers(client: TestClient) -> None:
    first = _texts(_chapter(client, 0))
    second = _chapter(client, 1)
    assert second["paragraphs"][0]["idx"] == 5
    # One more paragraph in chapter 1 (a line break also starts one), one fewer later.
    resp = _edit(
        client, 0, TITLES[0], [*first[:2], "她站了很久。\n  \n楼下有人按喇叭。", *first[2:]]
    )
    assert resp.status_code == 200, resp.text
    assert len(resp.json()["paragraphs"]) == 7
    moved = _chapter(client, 1)
    assert moved["paragraphs"][0]["idx"] == 7 and _texts(moved) == _texts(second)
    resp = _edit(client, 1, "第二章 十二点 · 改", _texts(moved)[1:])
    assert resp.status_code == 200, resp.text
    edited = resp.json()
    assert edited["chapters"][1]["title"] == "第二章 十二点 · 改"
    assert edited["chapters"][1]["para_count"] == second["chapters"][1]["para_count"] - 1
    assert _chapter(client, 2)["paragraphs"][0]["idx"] == edited["chapters"][2]["para_start"]
    assert len(_novel_product(client, "demo")["versions"]) == 3


def test_editing_chapter_3_leaves_what_reads_chapter_1_fresh(client: TestClient) -> None:
    state = _state(client)
    novel = _novel_product(client, "demo")
    chapter0 = novel["versions"][0]["payload"]["chapters"][0]

    async def downstream() -> str:
        p = await state.products.ensure(
            series_id="demo", owner_type="series", owner_id="demo", slot="test_chapter",
            component_id="test",
        )  # fmt: skip
        dep = DependencyDraft(novel["id"], slice_hash(chapter0), "chapter:0")
        await state.products.append(
            p.id, VersionDraft(payload={}, dependencies=[dep]), origin="run"
        )
        return p.id

    pid = client.portal.call(downstream)  # type: ignore[union-attr]
    last = _texts(_chapter(client, 2))
    assert _edit(client, 2, TITLES[2], [*last[:-1], "红伞底下，雨停了。"]).status_code == 200
    statuses = {
        p["component_id"]: p["status"]
        for p in client.get(f"{V1}/products", params={"owner": "series:demo"}).json()
    }
    assert statuses["test"] == "adopted"  # only read chapter 1
    assert statuses["extract"] == statuses["enrich"] == "stale"  # read the whole book
    assert client.get(f"{V1}/products/{pid}").json()["status"] == "adopted"


def test_suggestions_follow_an_edit(client: TestClient) -> None:
    """Paragraph numbers of the suggestions follow the text; an edited paragraph keeps its."""
    paras = _texts(_chapter(client, 0))
    paras[0] = paras[0].replace("加完班", "加完夜班")  # s1, s2 are on this paragraph
    resp = _edit(client, 0, TITLES[0], ["开头多了一段。", *paras])
    assert resp.status_code == 200, resp.text
    source = resp.json()["source"]["id"]
    data = client.get(f"{V1}/series/demo/suggestions").json()
    assert data["source_id"] == source
    assert [(s["id"], s["para_idx"]) for s in data["suggestions"]] == [
        ("s1", 1), ("s2", 1), ("s3", 2), ("s4", 4),
    ]  # fmt: skip
    assert all(s["anchor_span"]["source_id"] == source for s in data["suggestions"])


def test_edit_errors(client: TestClient) -> None:
    resp = _edit(client, 0, TITLES[0], ["  ", "\n"])
    assert (resp.status_code, _code(resp)) == (422, "chapter_empty")
    assert _code(_edit(client, 9, None, ["一段"])) == "chapter_not_found"
    resp = client.put(f"{V1}/series/nope/novel/chapters/0", json={"paragraphs": ["一段"]})
    assert _code(resp) == "series_not_found"
    sid = _new_series(client)
    resp = client.put(f"{V1}/series/{sid}/novel/chapters/0", json={"paragraphs": ["一段"]})
    assert (resp.status_code, _code(resp)) == (404, "novel_not_found")
    pid = _novel_product(client, "demo")["id"]
    assert client.post(f"{V1}/products/{pid}/lock").status_code == 200
    resp = _edit(client, 0, TITLES[0], ["一段"])
    assert (resp.status_code, _code(resp)) == (409, "product_locked")
